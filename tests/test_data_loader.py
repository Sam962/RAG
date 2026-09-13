import pytest

from src.data_loader import load_all_docs


def test_missing_data_dir_raises(tmp_path):
    with pytest.raises(FileNotFoundError, match="Data directory not found"):
        load_all_docs(str(tmp_path / "does-not-exist"))


def test_empty_data_dir_raises(tmp_path):
    with pytest.raises(ValueError, match="No documents loaded"):
        load_all_docs(str(tmp_path))


def test_loads_text_file(tmp_path):
    (tmp_path / "notes.txt").write_text("beam search is a decoding method.", encoding="utf-8")
    docs = load_all_docs(str(tmp_path))
    assert len(docs) == 1
    assert "beam search" in docs[0].page_content


def test_skips_unreadable_file_and_keeps_valid(tmp_path):
    (tmp_path / "good.txt").write_text("valid context", encoding="utf-8")
    (tmp_path / "bad.pdf").write_bytes(b"not a real pdf")
    docs = load_all_docs(str(tmp_path))
    assert len(docs) == 1
    assert docs[0].page_content == "valid context"


def test_all_files_unreadable_raises(tmp_path):
    (tmp_path / "bad.pdf").write_bytes(b"not a real pdf")
    with pytest.raises(ValueError, match="No documents loaded"):
        load_all_docs(str(tmp_path))


def test_loads_pptx_slides_and_notes(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches

    pptx_path = tmp_path / "deck.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(2))
    box.text_frame.text = "Beam search keeps the top k candidates."
    notes = slide.notes_slide.notes_text_frame
    notes.text = "Speaker note: greedy decoding is the k=1 special case."
    prs.save(str(pptx_path))

    docs = load_all_docs(str(tmp_path))
    assert len(docs) == 1
    assert "Beam search keeps the top k candidates." in docs[0].page_content
    assert "greedy decoding" in docs[0].page_content
    assert docs[0].metadata["page"] == 1
    assert docs[0].metadata["source"].endswith("deck.pptx")


def test_empty_pptx_is_skipped_and_raises_if_only_file(tmp_path):
    from pptx import Presentation

    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])
    prs.save(str(tmp_path / "empty.pptx"))
    with pytest.raises(ValueError, match="No documents loaded"):
        load_all_docs(str(tmp_path))

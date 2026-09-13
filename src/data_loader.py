"""Load PDF, TXT, CSV, and PPTX files into LangChain documents."""

from pathlib import Path
from typing import Any, Callable, List

from langchain_community.document_loaders import CSVLoader, PyPDFLoader, TextLoader
from langchain_core.documents import Document
from pptx import Presentation


def _load_pdf(path: Path) -> List[Any]:
    return PyPDFLoader(str(path)).load()


def _load_txt(path: Path) -> List[Any]:
    return TextLoader(str(path), encoding="utf-8").load()


def _load_csv(path: Path) -> List[Any]:
    return CSVLoader(str(path)).load()


def _load_pptx(path: Path) -> List[Document]:
    """One document per slide: shape text plus speaker notes."""
    presentation = Presentation(str(path))
    docs: List[Document] = []
    for slide_number, slide in enumerate(presentation.slides, start=1):
        parts: List[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                parts.append(shape.text_frame.text.strip())
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                parts.append(notes)
        content = "\n".join(parts)
        if not content:
            continue
        docs.append(
            Document(
                page_content=content,
                metadata={"source": str(path), "page": slide_number},
            )
        )
    if not docs:
        raise ValueError(f"No text found in PowerPoint: {path}")
    return docs


_LOADERS: List[tuple[str, str, Callable[[Path], List[Any]]]] = [
    ("PDF", "*.pdf", _load_pdf),
    ("TEXT", "*.txt", _load_txt),
    ("CSV", "*.csv", _load_csv),
    ("PPTX", "*.pptx", _load_pptx),
]


def load_all_docs(data_dir: str) -> List[Any]:
    """Load supported files from data_dir.

    Unreadable files are skipped. Raises if the directory is missing or nothing loads.
    """
    data_path = Path(data_dir).resolve()
    if not data_path.exists():
        raise FileNotFoundError(f"Data directory not found: {data_path}")

    print(f"[DEBUG] Data path : {data_path}")
    documents: List[Any] = []
    found = 0
    failed = 0

    for label, pattern, loader in _LOADERS:
        files = list(data_path.glob(f"**/{pattern}"))
        print(f"[DEBUG] Found {len(files)} {label} files: {[str(f) for f in files]}")
        for file_path in files:
            found += 1
            print(f"[DEBUG] Loading {label} {file_path}")
            try:
                loaded = loader(file_path)
                print(f"[DEBUG] loaded {len(loaded)} {label} docs from {file_path}")
                documents.extend(loaded)
            except Exception as e:
                failed += 1
                print(f"[ERROR] Failed to load {label} {file_path}: {e}")

    print(
        f"[INFO] Loaded {len(documents)} documents from {found} files "
        f"({failed} failed)"
    )
    if not documents:
        raise ValueError(
            f"No documents loaded from {data_path}. "
            f"Found {found} files, {failed} failed to read. "
            "Add PDF, TXT, CSV, or PPTX files under data/."
        )
    return documents
